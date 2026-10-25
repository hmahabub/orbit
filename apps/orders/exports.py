"""Downloadable documents: the supplier PO and the BOM (Excel and PDF) and
the cost sheet (Excel). The Excel files are laid out to be printed
(landscape, fitted to one page wide) and keep live formulas for amounts and
totals. A supplier PO that is not approved yet is marked as a draft."""

import datetime
import io

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .services import line_costing, load_bom

NAVY = '10233F'
THIN = Side(style='thin', color='9AA5B1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
LABEL_FILL = PatternFill('solid', fgColor='F3F6F9')
TOTAL_FILL = PatternFill('solid', fgColor='E3F3F3')
QTY, MONEY, RATE, CONS = '#,##0', '#,##0.00', '#,##0.0000', '0.000000'


def joined(values):
    """Distinct non-blank values, in first-seen order, as one string."""
    seen = []
    for value in values:
        if value and value not in seen:
            seen.append(value)
    return ', '.join(seen)


class Sheet:
    """Small helper around a worksheet: tracks the current row."""

    def __init__(self, title, widths):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = title[:31]
        self.row = 1
        self.cols = len(widths)
        for i, width in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = width
        self.ws.page_setup.orientation = 'landscape'
        self.ws.page_setup.paperSize = self.ws.PAPERSIZE_A4
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.print_options.horizontalCentered = True
        for side in ('left', 'right', 'top', 'bottom'):
            setattr(self.ws.page_margins, side, 0.4)

    def cell(self, col, value=None, *, bold=False, fmt=None, fill=None, align=None, box=True, color=None,
             size=None, wrap=False):
        c = self.ws.cell(row=self.row, column=col, value=value)
        if bold or color or size:
            c.font = Font(bold=bold, color=color, size=size)
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if box:
            c.border = BOX
        c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
        return c

    def merge(self, col_from, col_to, row_to=None):
        self.ws.merge_cells(start_row=self.row, start_column=col_from, end_row=row_to or self.row, end_column=col_to)

    def title(self, text, right_text, warning=None):
        mid = self.cols // 2
        if settings.COMPANY_LOGO.is_file():
            logo = Image(str(settings.COMPANY_LOGO))
            logo.height, logo.width = 50, round(50 * logo.width / logo.height)
            self.ws.add_image(logo, 'A1')
        self.cell(3, settings.COMPANY_NAME.upper(), bold=True, size=18, color=NAVY, box=False)
        self.merge(3, mid)
        self.cell(mid + 1, right_text, bold=True, size=14, color='FFFFFF', fill=HEAD_FILL, align='center', box=False)
        self.merge(mid + 1, self.cols)
        self.ws.row_dimensions[self.row].height = 42
        self.row += 1
        self.cell(1, settings.COMPANY_ADDRESS or None, box=False, color='445064')
        self.merge(1, self.cols)
        self.row += 2
        self.cell(1, text, bold=True, size=12, box=False)
        self.merge(1, self.cols)
        if warning:
            self.row += 1
            self.cell(1, warning, bold=True, color='C0392B', box=False)
            self.merge(1, self.cols)
        self.row += 2

    def info(self, pairs, spans):
        """Label/value pairs, two per row. spans = ((label cols), (value cols)) x 2."""
        for i in range(0, len(pairs), 2):
            for (label, value, *fmt), (lab, val) in zip(pairs[i:i + 2], spans):
                self.cell(lab[0], label, bold=True, fill=LABEL_FILL)
                for col in range(lab[0], val[1] + 1):
                    self.ws.cell(row=self.row, column=col).border = BOX
                if lab[1] > lab[0]:
                    self.merge(lab[0], lab[1])
                self.cell(val[0], value if value not in (None, '') else '—', fmt=fmt[0] if fmt else None, align='left')
                if val[1] > val[0]:
                    self.merge(val[0], val[1])
            self.row += 1
        self.row += 1

    def header(self, labels):
        """Column headings of the main table; repeated on every printed page."""
        for col, label in enumerate(labels, start=1):
            self.cell(col, label, bold=True, color='FFFFFF', fill=HEAD_FILL, align='center', wrap=True)
        self.ws.row_dimensions[self.row].height = 30
        self.ws.print_title_rows = f'{self.row}:{self.row}'
        self.row += 1

    def signatures(self, prepared_by, checked_by=''):
        """Prepared by = who created the document; Checked by = who approved it."""
        self.row += 3
        third = self.cols // 3
        for i, (name, label) in enumerate([(prepared_by, 'Prepared by'), (checked_by, 'Checked by'),
                                           ('', 'Authorized signature')]):
            first = 1 + i * third
            last = first + third - 2
            c = self.ws.cell(row=self.row, column=first, value=name)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal='center')
            self.ws.merge_cells(start_row=self.row, start_column=first, end_row=self.row, end_column=last)
            for col in range(first, last + 1):
                self.ws.cell(row=self.row, column=col).border = Border(bottom=Side(style='thin', color='333333'))
            c = self.ws.cell(row=self.row + 1, column=first, value=label)
            c.alignment = Alignment(horizontal='center')
            self.ws.merge_cells(start_row=self.row + 1, start_column=first, end_row=self.row + 1, end_column=last)

    def response(self, filename):
        out = io.BytesIO()
        self.wb.save(out)
        response = HttpResponse(
            out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        safe = ''.join(ch if ch.isalnum() or ch in '-_ ' else '-' for ch in filename).strip()
        response['Content-Disposition'] = f'attachment; filename="{safe}.xlsx"'
        return response


# --------------------------------------------------------------------------
# supplier PO
# --------------------------------------------------------------------------

def po_lines_with_detail(po):
    """The PO's lines, each carrying the detail of the BOM rows it was raised
    from: placement, body color and garment size."""
    lines = list(po.lines.prefetch_related('bom_rows__item__line__assortment'))
    for line in lines:
        rows = list(line.bom_rows.all())
        line.style = line.style or joined(row.item.line.style for row in rows)
        line.placement = joined(row.item.placement for row in rows)
        line.body_color = joined(row.body_color for row in rows) or ('All colors' if rows else '')
        line.garment_size = joined(row.size_label for row in rows)
    return lines


def po_details(po):
    """The PO header fields, in display order (pairs sit side by side)."""
    order = po.order
    return [
        ('Supplier', po.supplier), ('PO date', po.po_date),
        ('PI to', po.pi_to), ('PI to location', po.pi_to_location),
        ('PI no.', po.pi_no), ('LC no.', po.lc_no),
        ('Buyer', order.buyer), ('Buyer PO', order.po_number),
        ('Factory', order.factory), ('Factory location', po.factory_location),
        ('Destination', po.destination), ('Required delivery', po.delivery_date),
        ('ETD', po.etd), ('ETA', po.eta),
    ]


def po_approval(po):
    """(warning line, name for the Checked by slot: the approver and date)."""
    if po.is_approved:
        return None, f"{po.approver_name} · {po.approved_at:%d %b %Y}" if po.approved_at else po.approver_name
    return f'DRAFT — NOT APPROVED ({po.get_approval_display().lower()})', ''


PO_COLUMNS = ['#', 'Style', 'Item', 'Placement', 'Body color', 'Color combination', 'Size', 'Spec', 'Qty', 'Unit',
              'Price', 'Amount']


def supplier_po_xlsx(po, prepared_by):
    order = po.order
    sh = Sheet(po.po_no, [5, 21, 34, 18, 16, 18, 9, 12, 11, 8, 11, 14])
    warning, checked_by = po_approval(po)
    sh.title(f'Purchase Order · {po.type_label}', po.po_no, warning)
    sh.info([(label, value, 'dd mmm yyyy') if isinstance(value, datetime.date) else (label, value)
             for label, value in po_details(po)], spans=(((1, 2), (3, 5)), ((6, 7), (8, 12))))

    sh.header(PO_COLUMNS)
    first = sh.row
    for n, line in enumerate(po_lines_with_detail(po), start=1):
        for col, value in enumerate([n, line.style, line.description, line.placement, line.body_color, line.color,
                                     line.garment_size, line.spec], start=1):
            sh.cell(col, value, wrap=col in (2, 3, 4, 5, 6), align='center' if col == 1 else None)
        sh.cell(9, float(line.qty), fmt=QTY)
        sh.cell(10, line.unit)
        sh.cell(11, float(line.price) if line.price is not None else None, fmt=RATE)
        sh.cell(12, f'=I{sh.row}*K{sh.row}', fmt=MONEY)
        sh.row += 1
    last = sh.row - 1
    sh.cell(1, f'Total ({order.currency})', bold=True, fill=TOTAL_FILL, align='right')
    for col in range(2, 12):
        sh.cell(col, fill=TOTAL_FILL)
    sh.merge(1, 11)
    sh.cell(12, f'=SUM(L{first}:L{last})' if last >= first else 0, bold=True, fmt=MONEY, fill=TOTAL_FILL)
    sh.row += 1

    if po.notes:
        sh.row += 1
        sh.cell(1, f'Notes: {po.notes}', box=False, wrap=True)
        sh.merge(1, sh.cols)
        sh.ws.row_dimensions[sh.row].height = 15 * (po.notes.count('\n') + 1 + len(po.notes) // 120)
    sh.signatures(prepared_by, checked_by)
    return sh.response(f'{po.po_no} {po.supplier}')


class PdfDoc:
    """Landscape A4 PDF with the company letterhead, built with ReportLab."""

    def __init__(self, title, right_text, heading, warning=None):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Image as PdfImage
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

        self.colors, self.mm, self.Paragraph, self.Spacer, self.Table = colors, mm, Paragraph, Spacer, Table
        self.navy, self.grid = colors.HexColor('#' + NAVY), colors.HexColor('#9AA5B1')
        self.label_bg, self.total_bg = colors.HexColor('#F3F6F9'), colors.HexColor('#E3F3F3')
        self.base = ParagraphStyle('base', fontName='Helvetica', fontSize=8, leading=10)
        self.bold = ParagraphStyle('bold', parent=self.base, fontName='Helvetica-Bold')
        self.head = ParagraphStyle('head', parent=self.bold, textColor=colors.white, alignment=1)
        self.right = ParagraphStyle('right', parent=self.base, alignment=2)
        self.right_bold = ParagraphStyle('right_bold', parent=self.bold, alignment=2)
        self.center = ParagraphStyle('center', parent=self.base, alignment=1)
        self.center_bold = ParagraphStyle('center_bold', parent=self.bold, alignment=1)

        self.title, self.out = title, io.BytesIO()
        self.doc = SimpleDocTemplate(self.out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                                     topMargin=10 * mm, bottomMargin=12 * mm, title=title,
                                     author=settings.COMPANY_NAME)
        self.width = self.doc.width
        self.story = []

        # Letterhead: logo + company name + address on the left, a number box on the right.
        name = Paragraph(settings.COMPANY_NAME.upper(), ParagraphStyle(
            'co', fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=self.navy))
        address = Paragraph(self.esc(settings.COMPANY_ADDRESS), ParagraphStyle(
            'addr', parent=self.base, fontSize=8.5, textColor=colors.HexColor('#445064')))
        brand = [name, address]
        if settings.COMPANY_LOGO.is_file():
            from PIL import Image as PilImage
            with PilImage.open(settings.COMPANY_LOGO) as im:
                ratio = im.width / im.height
            # Size has to be given up front; ReportLab ignores it when set afterwards.
            logo = PdfImage(str(settings.COMPANY_LOGO), width=14 * mm * ratio, height=14 * mm)
            brand = Table([[logo, [name, address]]], colWidths=[logo.drawWidth + 4 * mm, None],
                          style=[('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0)])
        box = Table([[Paragraph(self.esc(right_text), ParagraphStyle(
            'box', fontName='Helvetica-Bold', fontSize=14, textColor=colors.white, alignment=1))]],
            colWidths=[70 * mm], style=[('BACKGROUND', (0, 0), (-1, -1), self.navy),
                                        ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 9)])
        self.add(Table([[brand, box]], colWidths=[self.width - 70 * mm, 70 * mm],
                       style=[('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0),
                              ('RIGHTPADDING', (0, 0), (-1, -1), 0)]), 4)
        self.add(Paragraph(self.esc(heading), ParagraphStyle('t', parent=self.bold, fontSize=11)), 1 if warning else 3)
        if warning:
            self.add(Paragraph(self.esc(warning), ParagraphStyle(
                'warn', parent=self.bold, fontSize=10, textColor=colors.HexColor('#C0392B'))), 3)

    @staticmethod
    def esc(text):
        return str(text).replace('&', '&amp;').replace('<', '&lt;').replace('\n', '<br/>')

    def p(self, value, style=None):
        if isinstance(value, datetime.date):
            value = value.strftime('%d %b %Y')
        return self.Paragraph('' if value in (None, '') else self.esc(value), style or self.base)

    def add(self, flowable, space_mm=0):
        self.story.append(flowable)
        if space_mm:
            self.story.append(self.Spacer(1, space_mm * self.mm))

    def details(self, pairs):
        """Label/value pairs, two per row."""
        rows = []
        for i in range(0, len(pairs), 2):
            row = []
            for label, value in pairs[i:i + 2]:
                row += [self.p(label, self.bold), self.p(value if value not in (None, '') else '—')]
            rows.append(row + [''] * (4 - len(row)))
        w = self.width / 2
        self.add(self.Table(rows, colWidths=[w * .32, w * .68, w * .32, w * .68], style=[
            ('GRID', (0, 0), (-1, -1), .5, self.grid), ('BACKGROUND', (0, 0), (0, -1), self.label_bg),
            ('BACKGROUND', (2, 0), (2, -1), self.label_bg), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]), 5)

    def table(self, data, shares, extra_style=(), total_row=False):
        """A grid with a navy heading row repeated on every page."""
        col_widths = [self.width * share / sum(shares) for share in shares]
        style = [('GRID', (0, 0), (-1, -1), .5, self.grid), ('BACKGROUND', (0, 0), (-1, 0), self.navy),
                 ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                 ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        if total_row:
            style += [('SPAN', (0, -1), (-2, -1)), ('BACKGROUND', (0, -1), (-1, -1), self.total_bg)]
        self.add(self.Table(data, colWidths=col_widths, repeatRows=1, style=style + list(extra_style)))

    def signatures(self, prepared_by, checked_by=''):
        """Prepared by = who created the document; Checked by = who approved it."""
        w = self.width
        self.add(self.Spacer(1, 16 * self.mm))
        line = self.colors.black
        self.add(self.Table(
            [[self.p(prepared_by, self.center_bold), '', self.p(checked_by, self.center_bold), '', ''],
             [self.p('Prepared by', self.center), '', self.p('Checked by', self.center), '',
              self.p('Authorized signature', self.center)]],
            colWidths=[w * .26, w * .11, w * .26, w * .11, w * .26],
            style=[('LINEBELOW', (0, 0), (0, 0), .7, line), ('LINEBELOW', (2, 0), (2, 0), .7, line),
                   ('LINEBELOW', (4, 0), (4, 0), .7, line)]))

    def response(self, filename):
        mm, colors, footer = self.mm, self.colors, self.title

        def page_number(canvas, doc):
            canvas.setFont('Helvetica', 7)
            canvas.setFillColor(colors.HexColor('#6B7686'))
            canvas.drawRightString(doc.pagesize[0] - 12 * mm, 6 * mm, f'{footer} · page {doc.page}')

        self.doc.build(self.story, onFirstPage=page_number, onLaterPages=page_number)
        response = HttpResponse(self.out.getvalue(), content_type='application/pdf')
        safe = ''.join(ch if ch.isalnum() or ch in '-_ ' else '-' for ch in filename).strip()
        response['Content-Disposition'] = f'attachment; filename="{safe}.pdf"'
        return response


def supplier_po_pdf(po, prepared_by):
    """The supplier PO as a PDF: same content and layout as the Excel one."""
    warning, checked_by = po_approval(po)
    pdf = PdfDoc(f'{po.po_no} {po.supplier}', po.po_no, f'Purchase Order · {po.type_label}', warning)
    pdf.details(po_details(po))
    p, right = pdf.p, pdf.right
    data = [[p(label, pdf.head) for label in PO_COLUMNS]]
    for n, line in enumerate(po_lines_with_detail(po), start=1):
        data.append([p(n), p(line.style), p(line.description), p(line.placement), p(line.body_color), p(line.color),
                     p(line.garment_size), p(line.spec), p(f'{line.qty:,.0f}', right), p(line.unit),
                     p(f'{line.price:,.4f}' if line.price is not None else '', right),
                     p(f'{line.amount:,.2f}', right)])
    data.append([p(f'Total ({po.order.currency})', pdf.right_bold)] + [''] * 10
                + [p(f'{po.total_amount:,.2f}', pdf.right_bold)])
    pdf.table(data, [3, 10, 18, 10, 9, 10, 5, 7, 7, 5, 7, 9], total_row=True)
    if po.notes:
        pdf.add(pdf.Spacer(1, 4 * pdf.mm))
        pdf.add(p(f'Notes: {po.notes}'))
    pdf.signatures(prepared_by, checked_by)
    return pdf.response(f'{po.po_no} {po.supplier}')


# --------------------------------------------------------------------------
# cost sheet
# --------------------------------------------------------------------------

def style_info(order, line, with_fob=True):
    """Order/style details; the BOM leaves out the FOB price."""
    pairs = [
        ('Buyer', order.buyer), ('Order', order.order_no),
        ('Factory', order.factory), ('Ship date', order.ship_date),
        ('Style', line.style), ('Item', line.description),
        ('Order qty (pcs)', line.qty),
    ]
    if with_fob:
        pairs.append((f'FOB / pc ({order.currency})', line.unit_price))
    return pairs


def style_header(sh, order, line, spans, pairs=None):
    """Order/style details (`pairs`, default: style_info with FOB) and the
    color x size breakdown at the top of a style's sheet. Returns the cells
    holding the order qty and FOB."""
    formats = {'Ship date': 'dd mmm yyyy', 'Order qty (pcs)': QTY}
    pairs = style_info(order, line) if pairs is None else pairs
    sh.info([(label, float(value) if label.startswith('FOB') else value,
              RATE if label.startswith('FOB') else formats.get(label)) for label, value in pairs],
            spans=spans)
    cells = (f'${get_column_letter(spans[0][1][0])}${sh.row - 2}', f'${get_column_letter(spans[1][1][0])}${sh.row - 2}')

    sizes = line.active_size_indexes
    sh.cell(1, 'Color / Size', bold=True, fill=LABEL_FILL)
    sh.cell(2, fill=LABEL_FILL)
    sh.merge(1, 2)
    for i, idx in enumerate(sizes):
        sh.cell(3 + i, line.size_label(idx), bold=True, fill=LABEL_FILL, align='center')
    sh.cell(3 + len(sizes), 'Total', bold=True, fill=LABEL_FILL, align='center')
    sh.row += 1
    first = sh.row
    for row in line.assortment.all():
        sh.cell(1, row.color)
        sh.cell(2)
        sh.merge(1, 2)
        for i, idx in enumerate(sizes):
            sh.cell(3 + i, row.qty_at(idx), fmt=QTY)
        end = get_column_letter(2 + len(sizes))
        sh.cell(3 + len(sizes), f'=SUM(C{sh.row}:{end}{sh.row})' if sizes else 0, bold=True, fmt=QTY)
        sh.row += 1
    if sh.row > first:
        sh.cell(1, 'Total', bold=True, fill=TOTAL_FILL)
        sh.cell(2, fill=TOTAL_FILL)
        sh.merge(1, 2)
        for i in range(len(sizes) + 1):
            letter = get_column_letter(3 + i)
            sh.cell(3 + i, f'=SUM({letter}{first}:{letter}{sh.row - 1})', bold=True, fmt=QTY, fill=TOTAL_FILL)
        sh.row += 1
    sh.row += 1
    return cells


def size_spec(row):
    return ' · '.join(v for v in [row.size_label if row.size_label != row.spec else '', row.spec] if v)


def cost_sheet_xlsx(order, line, prepared_by):
    items = load_bom(line)
    costing = line_costing(line, items)
    sh = Sheet(f'Cost {line.style}', [5, 12, 34, 18, 15, 18, 13, 11, 11, 12, 8, 10, 14])
    sh.title(f'Cost Sheet · {line.style}', order.po_number)
    qty_cell, fob_cell = style_header(sh, order, line, spans=(((1, 2), (3, 5)), ((6, 7), (8, 13))))

    # Materials. Item-level cells are merged down the item's rows, like the source sheet.
    sh.header(['Sl', 'Type', 'Item', 'Placement', 'Body color', 'Color combination', 'Size / spec', 'Order qty',
               'Cons. / pc', 'Total req.', 'Unit', 'Price', 'TTL amount'])
    first = sh.row
    for n, item in enumerate(items, start=1):
        top = sh.row
        for row in item.row_list:
            is_top = sh.row == top
            sh.cell(1, n if is_top else None, align='center')
            sh.cell(2, item.get_category_display() if is_top else None)
            sh.cell(3, item.name if is_top else None, wrap=True)
            sh.cell(4, item.placement if is_top else None, wrap=True)
            sh.cell(5, row.body_label)
            sh.cell(6, row.color_combo)
            sh.cell(7, size_spec(row))
            sh.cell(8, row.order_qty, fmt=QTY)
            sh.cell(9, float(row.consumption), fmt=CONS)
            sh.cell(10, f'=H{sh.row}*I{sh.row}', fmt=MONEY)
            sh.cell(11, item.unit if is_top else None)
            sh.cell(12, float(row.price) if row.price is not None else None, fmt=RATE)
            sh.cell(13, f'=J{sh.row}*L{sh.row}', fmt=MONEY)
            sh.row += 1
        if sh.row - top > 1:
            for col in (1, 2, 3, 4, 11):
                sh.ws.merge_cells(start_row=top, start_column=col, end_row=sh.row - 1, end_column=col)
    last = sh.row - 1
    sh.cell(1, f'Total material cost ({order.currency})', bold=True, fill=TOTAL_FILL, align='right')
    for col in range(2, 13):
        sh.cell(col, fill=TOTAL_FILL)
    sh.merge(1, 12)
    sh.cell(13, f'=SUM(M{first}:M{last})' if last >= first else 0, bold=True, fmt=MONEY, fill=TOTAL_FILL)
    material_cell = f'$M${sh.row}'
    sh.row += 2

    # Other costs (left) and summary (right), side by side.
    start = sh.row
    sh.cell(1, 'Other costs', bold=True, fill=LABEL_FILL)
    sh.cell(2, fill=LABEL_FILL)
    sh.cell(3, fill=LABEL_FILL)
    sh.merge(1, 3)
    sh.cell(4, 'Per pc', bold=True, fill=LABEL_FILL, align='right')
    sh.row += 1
    extras_first = sh.row
    for extra in costing['extras']:
        sh.cell(1, extra.name)
        sh.cell(2)
        sh.cell(3)
        sh.merge(1, 3)
        sh.cell(4, float(extra.cost_per_pc), fmt=RATE)
        sh.row += 1
    sh.cell(1, 'Total other costs / pc', bold=True, fill=TOTAL_FILL)
    sh.cell(2, fill=TOTAL_FILL)
    sh.cell(3, fill=TOTAL_FILL)
    sh.merge(1, 3)
    sh.cell(4, f'=SUM(D{extras_first}:D{sh.row - 1})' if sh.row > extras_first else 0, bold=True, fmt=RATE,
            fill=TOTAL_FILL)
    extras_cell = f'$D${sh.row}'
    left_end = sh.row

    sh.row = start
    r = start  # summary rows: label in F:H, value in I:J
    summary = [
        ('Material / pc', f'={material_cell}/{qty_cell}', RATE),
        ('Material / dozen', f'=I{r}*12', RATE),
        ('Other costs / pc', f'={extras_cell}', RATE),
        ('Total cost / pc', f'=I{r}+I{r + 2}', RATE),
        ('FOB / pc', f'={fob_cell}', RATE),
        ('Margin / pc', f'=I{r + 4}-I{r + 3}', RATE),
        ('Margin %', f'=IF(I{r + 4}=0,0,I{r + 5}/I{r + 4})', '0.0%'),
        ('Order margin', f'=I{r + 5}*{qty_cell}', MONEY),
    ]
    for label, formula, fmt in summary:
        strong = label in ('Total cost / pc', 'Margin / pc', 'Order margin')
        sh.cell(6, label, bold=True, fill=LABEL_FILL)
        sh.cell(7, fill=LABEL_FILL)
        sh.cell(8, fill=LABEL_FILL)
        sh.merge(6, 8)
        sh.cell(9, formula, bold=strong, fmt=fmt, fill=TOTAL_FILL if strong else None)
        sh.cell(10, fill=TOTAL_FILL if strong else None)
        sh.merge(9, 10)
        sh.row += 1
    sh.row = max(sh.row, left_end + 1)
    if costing['unpriced']:
        sh.cell(1, f"{costing['unpriced']} row(s) have no price yet and count as zero.", box=False, color='C0392B')
        sh.merge(1, sh.cols)
        sh.row += 1
    sh.signatures(prepared_by)
    return sh.response(f'Cost sheet {order.po_number} {line.style}')


# --------------------------------------------------------------------------
# BOM (material allocation) — no prices
# --------------------------------------------------------------------------

BOM_COLUMNS = ['Sl', 'Type', 'Item', 'Placement', 'Body color', 'Color combination', 'Size / spec', 'Order qty',
               'Cons. / pc', 'Total req.', 'Unit', 'Supplier']


def bom_xlsx(order, line, prepared_by):
    items = load_bom(line)
    sh = Sheet(f'BOM {line.style}', [5, 12, 34, 18, 15, 18, 13, 11, 11, 12, 8, 18])
    sh.title(f'Bill of Materials · {line.style}', order.po_number)
    style_header(sh, order, line, spans=(((1, 2), (3, 5)), ((6, 7), (8, 12))),
                 pairs=style_info(order, line, with_fob=False) + [('Materials', len(items))])
    sh.header(BOM_COLUMNS)
    for n, item in enumerate(items, start=1):
        top = sh.row
        for row in item.row_list:
            is_top = sh.row == top
            sh.cell(1, n if is_top else None, align='center')
            sh.cell(2, item.get_category_display() if is_top else None)
            sh.cell(3, item.name if is_top else None, wrap=True)
            sh.cell(4, item.placement if is_top else None, wrap=True)
            sh.cell(5, row.body_label)
            sh.cell(6, row.color_combo)
            sh.cell(7, size_spec(row))
            sh.cell(8, row.order_qty, fmt=QTY)
            sh.cell(9, float(row.consumption), fmt=CONS)
            sh.cell(10, f'=H{sh.row}*I{sh.row}', fmt=MONEY)
            sh.cell(11, item.unit if is_top else None)
            sh.cell(12, item.supplier if is_top else None, wrap=True)
            sh.row += 1
        if sh.row - top > 1:
            for col in (1, 2, 3, 4, 11, 12):
                sh.ws.merge_cells(start_row=top, start_column=col, end_row=sh.row - 1, end_column=col)
    if not items:
        sh.cell(1, 'No materials yet.', box=False)
        sh.row += 1
    sh.signatures(prepared_by)
    return sh.response(f'BOM {order.po_number} {line.style}')


def bom_pdf(order, line, prepared_by):
    items = load_bom(line)
    pdf = PdfDoc(f'BOM {order.po_number} {line.style}', order.po_number, f'Bill of Materials · {line.style}')
    pdf.details([(label, f'{value:,}' if isinstance(value, int) else value)
                 for label, value in style_info(order, line, with_fob=False) + [('Materials', len(items))]])
    p, right = pdf.p, pdf.right

    # Color x size breakdown.
    sizes = line.active_size_indexes
    grid = [[p('Color / Size', pdf.head)] + [p(line.size_label(i), pdf.head) for i in sizes] + [p('Total', pdf.head)]]
    for row in line.assortment.all():
        grid.append([p(row.color)] + [p(f'{row.qty_at(i):,}', right) for i in sizes] + [p(f'{row.total:,}', pdf.right_bold)])
    grid.append([p('Total', pdf.bold)] + [p(f'{line.size_total(i):,}', pdf.right_bold) for i in sizes]
                + [p(f'{line.assortment_total:,}', pdf.right_bold)])
    col = pdf.width * 0.45 / (len(sizes) + 2)
    pdf.add(pdf.Table(grid, colWidths=[col * 1.6] + [col] * len(sizes) + [col], hAlign='LEFT', style=[
        ('GRID', (0, 0), (-1, -1), .5, pdf.grid), ('BACKGROUND', (0, 0), (-1, 0), pdf.navy),
        ('BACKGROUND', (0, -1), (-1, -1), pdf.total_bg), ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]), 5)

    # Materials: item cells only on each material's first row, a heavier line between materials.
    data = [[p(label, pdf.head) for label in BOM_COLUMNS]]
    group_starts = []
    for n, item in enumerate(items, start=1):
        group_starts.append(len(data))
        for i, row in enumerate(item.row_list):
            first = i == 0
            data.append([p(n if first else ''), p(item.get_category_display() if first else ''),
                         p(item.name if first else ''), p(item.placement if first else ''), p(row.body_label),
                         p(row.color_combo), p(size_spec(row)), p(f'{row.order_qty:,}', right),
                         p(f'{row.consumption:.6f}'.rstrip('0').rstrip('.'), right),
                         p(f'{row.requirement:,.2f}', right), p(item.unit if first else ''),
                         p(item.supplier if first else '')])
    extra = [('LINEABOVE', (0, r), (-1, r), 1.2, pdf.navy) for r in group_starts[1:]]
    if not items:
        data.append([p('No materials yet.')] + [''] * 11)
    pdf.table(data, [3, 7, 18, 10, 9, 10, 8, 7, 7, 8, 5, 10], extra)
    pdf.signatures(prepared_by)
    return pdf.response(f'BOM {order.po_number} {line.style}')
