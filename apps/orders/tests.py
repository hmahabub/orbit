from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from .models import BOMItem, BOMRow, CostExtra, Order, SupplierPO


class WorkflowTests(TestCase):
    """Walks one order through all six tabs as a merchandiser and an admin."""

    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user('boss', password='x', is_staff=True)
        self.merch = User.objects.create_user('merch', password='x')
        self.client.force_login(self.merch)

    def xlsx_text(self, response):
        """Every cell of a downloaded workbook, as one string."""
        from io import BytesIO

        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        ws = load_workbook(BytesIO(response.content)).active
        return ' | '.join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)

    def detail(self, order, tab, **extra):
        params = {'tab': tab, **extra}
        return self.client.get(reverse('order_detail', args=[order.pk]), params)

    def create_order(self):
        self.client.post(reverse('order_create'), {
            'po_number': 'WC6916', 'buyer': 'N&I', 'factory': 'Afrah', 'ship_date': '2026-12-30',
            'currency': 'USD', 'status': 'running'})
        return Order.objects.get(po_number='WC6916')

    def add_line_and_assortment(self, order):
        self.client.post(reverse('line_add', args=[order.pk]), {
            'style': 'WC217G398', 'description': 'TRAVIS ZIP CARGO PANT', 'qty': 14040, 'unit_price': '4.60'})
        line = order.lines.get()
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_ASSORTMENT]))
        data = {'size_type': 'letter'}
        for n, (color, qtys) in enumerate([('OYSTER', [836, 1672, 1672, 836]), ('STEEL', [668, 1336, 1336, 668]),
                                           ('FERN CAMO', [836, 1672, 1672, 836])]):
            data[f'color-{n}'] = color
            for j, q in enumerate(qtys):
                data[f'q-{n}-{j}'] = q
        self.client.post(reverse('assortment_save', args=[order.pk, line.pk]), data)
        return order.lines.get()

    def test_tabs_unlock_in_order(self):
        order = self.create_order()
        self.assertEqual(order.stage, Order.STAGE_LINES)
        # Assortment is locked until a line exists and the step is completed.
        self.assertEqual(self.detail(order, 'assortment').context['tab'], 'lines')
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_ASSORTMENT]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_LINES)

        line = self.add_line_and_assortment(order)
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_ASSORTMENT)
        self.assertTrue(line.assortment_complete)

        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_BOM]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_BOM)

        # BOM can't be completed empty.
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_COSTING]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_BOM)

    def test_assortment_mismatch_blocks_bom(self):
        order = self.create_order()
        self.client.post(reverse('line_add', args=[order.pk]), {
            'style': 'S1', 'description': '', 'qty': 100, 'unit_price': '1'})
        line = order.lines.get()
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_ASSORTMENT]))
        self.client.post(reverse('assortment_save', args=[order.pk, line.pk]),
                         {'size_type': 'number', 'color-0': 'RED', 'q-0-0': 60})
        line.refresh_from_db()
        self.assertEqual(line.sizes[0], '30')
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_BOM]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_ASSORTMENT)

    def add_material(self, line, **fields):
        data = {'category': 'trims', 'placement': '', 'unit': 'Pcs', 'allocation': 'all',
                'consumption': '1', 'color_combo': '', 'spec': '', **fields}
        self.client.post(reverse('bom_item_add', args=[line.order_id, line.pk]), data)
        return BOMItem.objects.get(line=line, name=fields['name'])

    def test_bom_allocation_and_requirement_match_spreadsheet(self):
        order = self.create_order()
        line = self.add_line_and_assortment(order)
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_BOM]))
        twill = self.add_material(line, name='Twill', allocation='color', consumption='1.478070', spec='58"',
                                  category='fabric', unit='Yds')
        cord = self.add_material(line, name='Drawcord', allocation='size', color_combo='Black')
        self.add_material(line, name='Grommet', consumption='2')

        oyster = twill.rows.get(body_color='OYSTER')
        self.assertEqual(oyster.order_qty, 5016)
        self.assertEqual(oyster.color_combo, 'OYSTER')
        self.assertAlmostEqual(float(oyster.requirement), 7414.0, places=0)   # =G10/12*H10
        self.assertEqual([r.size_label for r in cord.rows.all()], ['S', 'M', 'L', 'XL'])
        self.assertEqual(cord.rows.get(size_idx=1).order_qty, 4680)
        grommet = BOMItem.objects.get(name='Grommet').rows.get()
        self.assertEqual(grommet.requirement, 14040 * 2)   # always per piece, no wastage

        # One save covers item cells and row cells; qty override; switching allocation rebuilds rows.
        self.client.post(reverse('bom_save', args=[order.pk, line.pk]), {
            f'i-{twill.pk}-name': 'Twill', f'i-{twill.pk}-allocation': 'color',
            f'i-{twill.pk}-unit': 'Yds', f'i-{twill.pk}-category': 'fabric', f'i-{twill.pk}-placement': 'Body',
            f'i-{cord.pk}-name': 'Drawcord', f'i-{cord.pk}-allocation': 'all',
            f'r-{oyster.pk}-cons': '1.478070', f'r-{oyster.pk}-qty': '5000', f'r-{oyster.pk}-src': '',
            f'r-{oyster.pk}-combo': 'OYSTER', f'r-{oyster.pk}-spec': '58"'})
        self.assertEqual(BOMRow.objects.get(pk=oyster.pk).order_qty, 5000)
        self.assertEqual(BOMItem.objects.get(pk=twill.pk).placement, 'Body')
        self.assertEqual(cord.rows.count(), 1)

        # Assortment change re-syncs rows but keeps what was typed.
        self.client.post(reverse('assortment_save', args=[order.pk, line.pk]),
                         {'size_type': 'letter', 'color-0': 'OYSTER', 'q-0-0': 14040})
        self.assertEqual(twill.rows.count(), 1)
        self.assertEqual(BOMRow.objects.get(pk=oyster.pk).consumption, Decimal('1.478070'))

    def test_total_requirement_drives_consumption(self):
        order = self.create_order()
        line = self.add_line_and_assortment(order)
        # Add row: a typed total is spread over the rows at one consumption / pc.
        fabric = self.add_material(line, name='Fabric', allocation='color', consumption='', total_req='20752')
        steel = fabric.rows.get(body_color='STEEL')
        self.assertAlmostEqual(float(steel.consumption), 20752 / 14040, places=8)
        self.assertAlmostEqual(float(sum(r.requirement for r in fabric.rows.all())), 20752, places=4)

        # Sheet: typing Total req. on a row sets its consumption = req / qty.
        oyster = fabric.rows.get(body_color='OYSTER')
        post = {f'i-{fabric.pk}-name': 'Fabric', f'i-{fabric.pk}-allocation': 'color'}
        for r in fabric.rows.all():
            post.update({f'r-{r.pk}-cons': '9', f'r-{r.pk}-req': '1', f'r-{r.pk}-qty': '', f'r-{r.pk}-src': '',
                         f'r-{r.pk}-combo': r.color_combo, f'r-{r.pk}-spec': ''})
        post.update({f'r-{oyster.pk}-req': '7414', f'r-{oyster.pk}-src': 'req',
                     f'r-{steel.pk}-cons': '1.5', f'r-{steel.pk}-src': 'cons'})
        self.client.post(reverse('bom_save', args=[order.pk, line.pk]), post)
        oyster.refresh_from_db()
        self.assertAlmostEqual(float(oyster.requirement), 7414, places=4)
        self.assertAlmostEqual(float(oyster.consumption), 7414 / 5016, places=8)
        self.assertEqual(BOMRow.objects.get(pk=steel.pk).consumption, Decimal('1.5'))
        # Untouched row keeps its stored consumption (src empty ignores the posted '9').
        camo = fabric.rows.get(body_color='FERN CAMO')
        self.assertAlmostEqual(float(camo.consumption), 20752 / 14040, places=8)

    def test_costing_and_supplier_po_workflow(self):
        """A merchandiser (not staff) costs the order and raises POs; staff approve them."""
        order = self.create_order()
        self.add_line_and_assortment(order)
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_BOM]))
        line = order.lines.get()
        twill = self.add_material(line, name='Twill', category='fabric', unit='Yds', allocation='color',
                                  consumption='1.478070', spec='58"', placement='Main body')
        zipper = self.add_material(line, name='Zipper', category='trims')
        label = self.add_material(line, name='Care label', category='labels')
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_COSTING]))

        # Everyone can cost now.
        self.assertEqual(self.detail(order, 'costing').context['tab'], 'costing')
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_PO]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_COSTING)  # unpriced rows block it
        prices = {f'r-{r.pk}-price': '1.85' for r in BOMRow.objects.all()}
        self.client.post(reverse('costing_save', args=[order.pk, line.pk]), prices)
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_PO]))
        order.refresh_from_db()
        self.assertEqual(order.stage, Order.STAGE_PO)

        # Suppliers are chosen on the PO tab; one PO per supplier and material type, pending approval.
        assign = reverse('po_assign', args=[order.pk])
        all_po_ids = lambda: list(SupplierPO.objects.values_list('pk', flat=True))  # noqa: E731
        pending = lambda: [i.name for i in self.detail(order, 'po').context['pending']]  # noqa: E731
        self.client.post(assign, {f's-{twill.pk}': 'Roundstone', f's-{zipper.pk}': 'Roundstone',
                                  f's-{label.pk}': '', 'action': 'generate'})
        self.assertEqual(sorted(SupplierPO.objects.values_list('supplier', 'material_type', 'approval')),
                         [('Roundstone', 'fabric', 'pending'), ('Roundstone', 'trims', 'pending')])
        po = SupplierPO.objects.get(material_type='fabric')
        trims_po = SupplierPO.objects.get(material_type='trims')
        self.assertEqual(po.created_by, self.merch)
        self.assertEqual(po.lines.get(color='OYSTER').qty, 7414)
        self.assertEqual(pending(), ['Care label'])
        self.client.post(assign, {'action': 'generate'})
        self.assertEqual(SupplierPO.objects.count(), 2)

        # Move / take off / split / merge — all by the raiser.
        steel = po.lines.get(color='STEEL')
        self.client.post(reverse('po_move', args=[order.pk, po.pk]), {'line': [steel.pk], 'target': trims_po.pk})
        trims_po.refresh_from_db()
        self.assertEqual((trims_po.lines.count(), trims_po.type_label), (2, 'Mixed'))
        self.client.post(reverse('po_move', args=[order.pk, po.pk]),
                         {'line': [po.lines.get(color='FERN CAMO').pk], 'target': 'pending'})
        self.assertEqual(pending(), ['Twill', 'Care label'])
        self.client.post(reverse('po_move', args=[order.pk, trims_po.pk]), {'line': [steel.pk], 'target': 'new'})
        self.assertEqual(SupplierPO.objects.count(), 3)
        self.client.post(reverse('po_merge', args=[order.pk]), {'po': all_po_ids()})
        po = SupplierPO.objects.get()
        self.assertEqual(po.lines.count(), 3)
        self.assertRegex(order.order_no, r'^ORD-\d{2}-\d{5}$')
        self.assertRegex(po.po_no, r'^SPO-\d{2}-\d{5}$')

        # Details: an unapproved PO stays a draft whatever status is chosen.
        prefix = f'po{po.pk}-'
        details = {prefix + 'supplier': 'Roundstone', prefix + 'material_type': 'fabric', prefix + 'status': 'issued',
                   prefix + 'po_date': '2026-10-10', prefix + 'pi_no': 'PI-77', prefix + 'pi_to': 'Orbit Sourcing Ltd',
                   prefix + 'pi_to_location': 'Dhaka', prefix + 'factory_location': 'Gazipur',
                   prefix + 'destination': 'Chattogram port', prefix + 'etd': '2026-11-20'}
        self.client.post(reverse('po_update', args=[order.pk, po.pk]), details)
        po.refresh_from_db()
        self.assertEqual((po.status, po.pi_to, po.destination), ('draft', 'Orbit Sourcing Ltd', 'Chattogram port'))

        # Downloads of an unapproved PO say so.
        letterhead = self.settings(COMPANY_NAME='Acme Textiles', COMPANY_ADDRESS='12 Test Road, Dhaka')
        letterhead.enable()
        self.addCleanup(letterhead.disable)
        text = self.xlsx_text(self.client.get(reverse('po_excel', args=[po.pk])))
        for expected in ['ACME TEXTILES', '12 Test Road, Dhaka', po.po_no, 'PI-77', 'Gazipur', 'Main body',
                         'STEEL', '58"', 'merch', 'Prepared by', 'DRAFT — NOT APPROVED']:
            self.assertIn(expected, text, expected)
        self.assertNotIn('Ship date', text)
        pdf = self.client.get(reverse('po_pdf', args=[po.pk]))
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        # ...and the on-screen view shows the same data without downloading.
        view = self.client.get(reverse('po_view', args=[po.pk]))
        for expected in [po.po_no, 'PI-77', 'Gazipur', 'Chattogram port', 'Main body', 'STEEL', '58&quot;',
                         'DRAFT — NOT APPROVED', 'Prepared by', 'merch']:
            self.assertContains(view, expected, msg_prefix=expected)

        # Another non-staff user can see and download it, but not change it or approve it.
        other = get_user_model().objects.create_user('other', password='x')
        self.client.force_login(other)
        self.assertFalse(self.detail(order, 'po').context['pos'][0][0].can_edit)
        self.assertEqual(self.client.get(reverse('po_excel', args=[po.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('po_view', args=[po.pk])).status_code, 200)
        for name in ['po_move', 'po_update', 'po_lines_save', 'po_delete', 'po_resubmit', 'po_approve', 'po_reject',
                     'po_reopen']:
            self.assertEqual(self.client.post(reverse(name, args=[order.pk, po.pk])).status_code, 403, name)
        self.assertEqual(self.client.post(reverse('po_merge', args=[order.pk]), {'po': [po.pk, po.pk]}).status_code, 403)
        self.assertEqual(self.client.get(reverse('approvals')).status_code, 403)
        self.assertEqual(self.client.post(reverse('order_delete', args=[order.pk])).status_code, 403)
        # The raiser can't approve their own PO either, unless they are staff.
        self.client.force_login(self.merch)
        self.assertEqual(self.client.post(reverse('po_approve', args=[order.pk, po.pk])).status_code, 403)

        # Staff: sends it back with a reason; the raiser corrects and resubmits.
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse('approvals')), po.po_no)
        self.client.post(reverse('po_reject', args=[order.pk, po.pk]), {'note': 'Check fabric price'})
        po.refresh_from_db()
        self.assertEqual((po.approval, po.approval_note), ('rejected', 'Check fabric price'))
        self.client.force_login(self.merch)
        self.client.post(reverse('po_resubmit', args=[order.pk, po.pk]))
        po.refresh_from_db()
        self.assertEqual(po.approval, 'pending')

        # Staff approves: lines lock for everyone; details stay editable by the raiser, supplier is fixed.
        self.client.force_login(self.admin)
        self.client.post(reverse('po_approve', args=[order.pk, po.pk]))
        po.refresh_from_db()
        self.assertEqual((po.approval, po.approved_by, po.status), ('approved', self.admin, 'issued'))
        self.assertNotIn(po, self.client.get(reverse('approvals')).context['pos'])
        self.assertEqual(self.client.post(reverse('po_lines_save', args=[order.pk, po.pk])).status_code, 403)
        self.client.force_login(self.merch)
        self.assertEqual(self.client.post(reverse('po_delete', args=[order.pk, po.pk])).status_code, 403)
        self.client.post(reverse('po_update', args=[order.pk, po.pk]), {**details, prefix + 'supplier': 'Someone else',
                                                                        prefix + 'status': 'draft'})
        po.refresh_from_db()
        self.assertEqual((po.status, po.supplier), ('issued', 'Roundstone'))  # approved: never back to draft
        self.client.post(reverse('po_update', args=[order.pk, po.pk]), {**details, prefix + 'status': 'closed'})
        po.refresh_from_db()
        self.assertEqual(po.status, 'closed')
        text = self.xlsx_text(self.client.get(reverse('po_excel', args=[po.pk])))
        self.assertNotIn('NOT APPROVED', text)
        self.assertIn('boss', text)  # approver under Authorized signature

        # Staff can reopen it for changes; it then needs approving again.
        self.client.force_login(self.admin)
        self.client.post(reverse('po_reopen', args=[order.pk, po.pk]))
        po.refresh_from_db()
        self.assertEqual((po.approval, po.status), ('pending', 'draft'))

        # Cost sheet and BOM downloads, for anyone.
        self.client.force_login(self.merch)
        text = self.xlsx_text(self.client.get(reverse('costing_excel', args=[order.pk, line.pk])))
        for expected in ['Cost Sheet', 'Twill', 'Total material cost', 'Margin / pc', order.order_no]:
            self.assertIn(expected, text, expected)
        text = self.xlsx_text(self.client.get(reverse('bom_download', args=[order.pk, line.pk, 'excel'])))
        for expected in ['Bill of Materials', 'Twill', 'Main body', 'OYSTER', '58"', 'Roundstone', 'Supplier',
                         'merch']:
            self.assertIn(expected, text, expected)
        self.assertNotIn('Price', text)
        self.assertNotIn('FOB', text)
        pdf = self.client.get(reverse('bom_download', args=[order.pk, line.pk, 'pdf']))
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        self.assertEqual(self.client.get(reverse('bom_download', args=[order.pk, line.pk, 'doc'])).status_code, 404)

        # The order's creator may delete it.
        self.assertEqual(self.client.post(reverse('order_delete', args=[order.pk])).status_code, 302)
        self.assertFalse(Order.objects.filter(pk=order.pk).exists())

    def test_only_the_creator_or_staff_can_change_an_order(self):
        order = self.create_order()            # created by merch
        line = self.add_line_and_assortment(order)
        self.client.post(reverse('advance', args=[order.pk, Order.STAGE_BOM]))
        # Supplier can be given when the material is added to the BOM.
        twill = self.add_material(line, name='Twill', supplier='Roundstone', unit='Yds')
        self.assertEqual(twill.supplier, 'Roundstone')
        self.client.post(reverse('bom_save', args=[order.pk, line.pk]), {
            f'i-{twill.pk}-name': 'Twill', f'i-{twill.pk}-supplier': 'Mashfiq', f'i-{twill.pk}-unit': 'Yds'})
        twill.refresh_from_db()
        self.assertEqual(twill.supplier, 'Mashfiq')
        row = twill.rows.get()
        extra = CostExtra.objects.create(line=line, name='CM', cost_per_pc=1)
        po = SupplierPO.objects.create(order=order, supplier='Mashfiq', created_by=self.merch)

        other = get_user_model().objects.create_user('other', password='x')
        self.client.force_login(other)
        # Can look, read-only.
        resp = self.detail(order, 'bom')
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.context['can_edit'])
        self.assertContains(resp, 'View only')
        self.assertEqual(self.client.get(reverse('bom_download', args=[order.pk, line.pk, 'excel'])).status_code, 200)
        # Can't change anything.
        blocked = [
            ('order_update', [order.pk]), ('order_delete', [order.pk]), ('advance', [order.pk, Order.STAGE_COSTING]),
            ('line_add', [order.pk]), ('line_update', [order.pk, line.pk]), ('line_delete', [order.pk, line.pk]),
            ('assortment_save', [order.pk, line.pk]), ('bom_item_add', [order.pk, line.pk]),
            ('bom_save', [order.pk, line.pk]), ('bom_item_delete', [order.pk, twill.pk]),
            ('costing_save', [order.pk, line.pk]), ('extra_add', [order.pk, line.pk]),
            ('extra_delete', [order.pk, extra.pk]), ('po_assign', [order.pk]), ('po_merge', [order.pk]),
            ('po_move', [order.pk, po.pk]), ('po_update', [order.pk, po.pk]), ('po_lines_save', [order.pk, po.pk]),
            ('po_delete', [order.pk, po.pk]), ('po_resubmit', [order.pk, po.pk]),
        ]
        for name, args in blocked:
            self.assertEqual(self.client.post(reverse(name, args=args)).status_code, 403, name)
        self.assertTrue(BOMItem.objects.filter(pk=twill.pk).exists())
        self.assertEqual(BOMRow.objects.get(pk=row.pk).consumption, row.consumption)

        # Staff can change it.
        self.client.force_login(self.admin)
        self.assertTrue(self.detail(order, 'bom').context['can_edit'])
        self.client.post(reverse('line_update', args=[order.pk, line.pk]), {
            f'ln{line.pk}-style': 'NEW STYLE', f'ln{line.pk}-description': '', f'ln{line.pk}-qty': line.qty,
            f'ln{line.pk}-unit_price': '5'})
        line.refresh_from_db()
        self.assertEqual(line.style, 'NEW STYLE')

    def test_pages_render(self):
        from django.core.management import call_command
        call_command('seed_demo', stdout=open('nul' if __import__('os').name == 'nt' else '/dev/null', 'w'))
        self.client.force_login(self.admin)
        order = Order.objects.get(po_number='WC6916')
        for tab in ['order', 'lines', 'assortment', 'bom', 'costing', 'po']:
            resp = self.detail(order, tab)
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(resp.context['tab'], tab)
        for name in ['home', 'po_list', 'order_create']:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        self.assertEqual(self.client.get(reverse('home'), {'status': 'all', 'q': 'TRAVIS'}).status_code, 200)

    def test_reports_and_downloads(self):
        import os
        from io import BytesIO

        from django.core.management import call_command
        from openpyxl import load_workbook

        from .reports import REPORTS
        call_command('seed_demo', stdout=open(os.devnull, 'w'))

        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('report_list')).status_code, 200)
        for key in REPORTS:
            url = reverse('report_detail', args=[key])
            self.assertEqual(self.client.get(url, {'status': 'running', 'buyer': 'N&I'}).status_code, 200, key)
            self.assertIn('text/csv', self.client.get(url, {'download': 'csv'})['Content-Type'])
            wb = load_workbook(BytesIO(self.client.get(url, {'download': 'xlsx'}).content))
            self.assertGreater(wb.active.max_row, 4, key)

        table = self.client.get(reverse('report_detail', args=['orders'])).context['table']
        self.assertEqual(table.totals[-2:], [56160, Decimal('258336.00')])
        # Ship-date filter narrows the selection.
        table = self.client.get(reverse('report_detail', args=['orders']), {'ship_from': '2026-06-01'}).context['table']
        self.assertEqual([row[1] for row in table.rows], ['WC6916'])

        # Every user sees every report, prices included.
        self.client.force_login(self.merch)
        for key in REPORTS:
            self.assertEqual(self.client.get(reverse('report_detail', args=[key])).status_code, 200, key)
        table = self.client.get(reverse('report_detail', args=['materials'])).context['table']
        self.assertTrue({'Amount', 'Supplier', 'PO status'} <= {c.label for c in table.columns})
        self.assertContains(self.client.get(reverse('report_list')), 'Costing &amp; margin')
        table = self.client.get(reverse('report_detail', args=['po-register'])).context['table']
        self.assertIn('Approval', [c.label for c in table.columns])
