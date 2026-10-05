"""Reports: each one turns the filtered orders into a plain table (columns +
rows + totals) that is shown on screen and downloaded as Excel or CSV from
the same data. Columns and reports marked admin-only carry prices/costs and
are left out for everyone else."""

import csv
import datetime
import io
from collections import OrderedDict
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

from .models import MATERIAL_TYPES, ZERO, Order, SupplierPO
from .services import line_costing, load_bom

STAGE_LABELS = {1: 'Order', 2: 'Line Items', 3: 'Assortment', 4: 'BOM', 5: 'Costing', 6: 'Supplier PO'}


class Col:
    """kind: text | int | qty (2 dp) | money (2 dp) | rate (4 dp) | pct | date"""

    def __init__(self, label, kind='text', total=False, admin=False):
        self.label, self.kind, self.total, self.admin = label, kind, total, admin

    @property
    def numeric(self):
        return self.kind in ('int', 'qty', 'money', 'rate', 'pct')


PLACES = {'int': 0, 'qty': 2, 'money': 2, 'rate': 4, 'pct': 1}


def show(value, kind):
    """A cell as display text."""
    if value is None or value == '':
        return ''
    if kind == 'date':
        return value.strftime('%d %b %Y')
    if kind in PLACES:
        return f'{value:,.{PLACES[kind]}f}' + ('%' if kind == 'pct' else '')
    return str(value)


class Table:
    def __init__(self, columns, rows, chart=None):
        self.columns, self.rows, self.chart = columns, rows, chart

    def for_user(self, admin):
        """Drop admin-only columns for non-admins."""
        if admin:
            return self
        keep = [i for i, col in enumerate(self.columns) if not col.admin]
        chart = self.chart
        if chart and (chart[0] not in keep or chart[1] not in keep):
            chart = None
        elif chart:
            chart = (keep.index(chart[0]), keep.index(chart[1]))
        return Table([self.columns[i] for i in keep], [[row[i] for i in keep] for row in self.rows], chart)

    @property
    def totals(self):
        if not self.rows or not any(col.total for col in self.columns):
            return None
        out = []
        for i, col in enumerate(self.columns):
            if col.total:
                out.append(sum((row[i] or 0 for row in self.rows), ZERO if col.kind != 'int' else 0))
            else:
                out.append('Total' if i == 0 else None)
        return out

    def _show(self, row):
        return [(show(value, col.kind), col.numeric) for value, col in zip(row, self.columns)]

    @property
    def display_rows(self):
        """Rows as (text, right-align?) pairs for the HTML table."""
        return [self._show(row) for row in self.rows]

    @property
    def display_totals(self):
        totals = self.totals
        return self._show(totals) if totals else None

    @property
    def bars(self):
        """Top rows of the chart column as (label, value text, % of the largest, negative?)."""
        if not self.chart or not self.rows:
            return []
        label_i, value_i = self.chart
        pairs = [(str(row[label_i]), row[value_i] or 0) for row in self.rows]
        pairs = sorted(pairs, key=lambda p: p[1], reverse=True)[:12]
        top = max((abs(v) for _, v in pairs), default=0)
        kind = self.columns[value_i].kind
        return [(label, show(value, kind), float(abs(value)) / float(top) * 100 if top else 0, value < 0)
                for label, value in pairs]

    @property
    def chart_label(self):
        return self.columns[self.chart[1]].label if self.chart else ''


# --------------------------------------------------------------------------
# filters
# --------------------------------------------------------------------------

def parse_date(value):
    try:
        return datetime.date.fromisoformat(value) if value else None
    except ValueError:
        return None


def read_filters(params):
    return {
        'status': params.get('status', '') if params.get('status') in dict(Order.STATUS_CHOICES) else '',
        'buyer': params.get('buyer', '').strip(),
        'ship_from': parse_date(params.get('ship_from')),
        'ship_to': parse_date(params.get('ship_to')),
    }


def filtered_orders(filters):
    orders = Order.objects.prefetch_related('lines__assortment').order_by('ship_date', 'id')
    if filters['status']:
        orders = orders.filter(status=filters['status'])
    if filters['buyer']:
        orders = orders.filter(buyer__iexact=filters['buyer'])
    if filters['ship_from']:
        orders = orders.filter(ship_date__gte=filters['ship_from'])
    if filters['ship_to']:
        orders = orders.filter(ship_date__lte=filters['ship_to'])
    return list(orders)


# --------------------------------------------------------------------------
# the reports
# --------------------------------------------------------------------------

def order_register(orders):
    cols = [Col('Order'), Col('Buyer PO'), Col('Buyer'), Col('Factory'), Col('Styles'), Col('Ship date', 'date'),
            Col('Status'), Col('Stage'), Col('Qty (pcs)', 'int', total=True), Col('Value', 'money', total=True)]
    rows = [[o.order_no, o.po_number, o.buyer, o.factory, o.descriptions, o.ship_date, o.get_status_display(),
             STAGE_LABELS.get(o.stage, ''), o.total_qty, o.total_amount] for o in orders]
    return Table(cols, rows)


def buyer_summary(orders):
    groups = OrderedDict()
    for o in orders:
        g = groups.setdefault(o.buyer, {'orders': 0, 'running': 0, 'qty': 0, 'value': ZERO})
        g['orders'] += 1
        g['running'] += o.status == 'running'
        g['qty'] += o.total_qty
        g['value'] += o.total_amount
    cols = [Col('Buyer'), Col('Orders', 'int', total=True), Col('Running', 'int', total=True),
            Col('Qty (pcs)', 'int', total=True), Col('Value', 'money', total=True), Col('Avg price / pc', 'rate')]
    rows = [[buyer, g['orders'], g['running'], g['qty'], g['value'], g['value'] / g['qty'] if g['qty'] else ZERO]
            for buyer, g in groups.items()]
    rows.sort(key=lambda r: r[4], reverse=True)
    return Table(cols, rows, chart=(0, 4))


def shipment_plan(orders):
    today = timezone.localdate()
    groups = OrderedDict()
    for o in sorted(orders, key=lambda o: (o.ship_date is None, o.ship_date or today)):
        key = o.ship_date.strftime('%Y-%m') if o.ship_date else 'No ship date'
        g = groups.setdefault(key, {'orders': 0, 'qty': 0, 'value': ZERO, 'overdue': 0})
        g['orders'] += 1
        g['qty'] += o.total_qty
        g['value'] += o.total_amount
        g['overdue'] += o.ship_status == 'red'
    cols = [Col('Ship month'), Col('Orders', 'int', total=True), Col('Qty (pcs)', 'int', total=True),
            Col('Value', 'money', total=True), Col('Past ship date', 'int', total=True)]
    rows = [[month, g['orders'], g['qty'], g['value'], g['overdue']] for month, g in groups.items()]
    return Table(cols, rows, chart=(0, 2))


def material_requirement(orders):
    types = dict(MATERIAL_TYPES)
    cols = [Col('Buyer PO'), Col('Style'), Col('Type'), Col('Item'), Col('Placement'), Col('Rows', 'int'),
            Col('Total req.', 'qty'), Col('Unit'), Col('Supplier'), Col('PO status'),
            Col('Amount', 'money', total=True, admin=True)]
    rows = []
    for o in orders:
        for line in o.lines.all():
            for item in load_bom(line):
                on_po = sum(1 for row in item.row_list if row.po_lines.all())
                status = 'Not ordered' if not on_po else ('On PO' if on_po == len(item.row_list) else 'Partly on PO')
                rows.append([o.po_number, line.style, types.get(item.category, ''), item.name, item.placement,
                             len(item.row_list), sum((r.requirement for r in item.row_list), ZERO), item.unit,
                             item.supplier, status, sum((r.amount for r in item.row_list), ZERO)])
    return Table(cols, rows)


def material_cost_by_type(orders):
    cols = [Col('Buyer PO'), Col('Buyer'), Col('Qty (pcs)', 'int', total=True)]
    cols += [Col(label, 'money', total=True) for _, label in MATERIAL_TYPES]
    cols += [Col('Total material', 'money', total=True), Col('Material / pc', 'rate')]
    rows = []
    for o in orders:
        by_type = {key: ZERO for key, _ in MATERIAL_TYPES}
        for line in o.lines.all():
            for item in load_bom(line):
                by_type[item.category] = by_type.get(item.category, ZERO) + sum((r.amount for r in item.row_list), ZERO)
        total = sum(by_type.values(), ZERO)
        if total:
            qty = o.total_qty
            rows.append([o.po_number, o.buyer, qty] + [by_type[key] for key, _ in MATERIAL_TYPES]
                        + [total, total / qty if qty else ZERO])
    return Table(cols, rows, chart=(0, len(cols) - 2))


def profitability(orders):
    cols = [Col('Buyer PO'), Col('Buyer'), Col('Style'), Col('Qty (pcs)', 'int', total=True), Col('FOB / pc', 'rate'),
            Col('Material / pc', 'rate'), Col('Other / pc', 'rate'), Col('Total cost / pc', 'rate'),
            Col('Margin / pc', 'rate'), Col('Margin %', 'pct'), Col('Order value', 'money', total=True),
            Col('Order margin', 'money', total=True), Col('Costing')]
    rows = []
    for o in orders:
        for line in o.lines.all():
            c = line_costing(line)
            state = 'No BOM' if not c['row_count'] else (f"{c['unpriced']} price(s) missing" if c['unpriced'] else 'Complete')
            rows.append([o.po_number, o.buyer, line.style, line.qty, c['fob'], c['material_pc'], c['extras_pc'],
                         c['total_pc'], c['margin_pc'], c['margin_pct'], line.amount, c['margin_total'], state])
    return Table(cols, rows, chart=(2, 11))


def _pos_for(orders):
    return (SupplierPO.objects.filter(order__in=orders).select_related('order').prefetch_related('lines')
            .order_by('supplier', 'id'))


def supplier_summary(orders):
    groups = OrderedDict()
    for po in _pos_for(orders):
        g = groups.setdefault(po.supplier, {'pos': 0, 'lines': 0, 'open': 0, 'received': 0, 'types': set(),
                                            'amount': ZERO})
        g['pos'] += 1
        g['lines'] += len(po.lines.all())
        g['received' if po.status == 'received' else 'open'] += 1
        g['types'].add(po.type_label)
        g['amount'] += po.total_amount
    cols = [Col('Supplier'), Col('Types'), Col('POs', 'int', total=True), Col('Lines', 'int', total=True),
            Col('Open POs', 'int', total=True), Col('Received', 'int', total=True),
            Col('PO amount', 'money', total=True, admin=True)]
    rows = [[supplier, ', '.join(sorted(g['types'])), g['pos'], g['lines'], g['open'], g['received'], g['amount']]
            for supplier, g in groups.items()]
    return Table(cols, rows, chart=(0, 6))


def po_register(orders):
    cols = [Col('PO no.'), Col('Supplier'), Col('Type'), Col('Buyer PO'), Col('Buyer'), Col('PO date', 'date'),
            Col('Delivery', 'date'), Col('PI no.'), Col('LC no.'), Col('ETD', 'date'), Col('ETA', 'date'),
            Col('Status'), Col('Lines', 'int', total=True), Col('Amount', 'money', total=True, admin=True)]
    rows = [[po.po_no, po.supplier, po.type_label, po.order.po_number, po.order.buyer, po.po_date, po.delivery_date,
             po.pi_no, po.lc_no, po.etd, po.eta, po.get_status_display(), len(po.lines.all()), po.total_amount]
            for po in _pos_for(orders)]
    return Table(cols, rows)


# key -> (title, what it answers, icon, admin only, builder)
REPORTS = OrderedDict([
    ('orders', ('Order register', 'Every order with qty, value, ship date and how far it has progressed.',
                'bi-clipboard-data', False, order_register)),
    ('buyers', ('Buyer summary', 'Orders, pieces and value per buyer.', 'bi-people', False, buyer_summary)),
    ('shipments', ('Shipment plan', 'Orders, pieces and value shipping each month, and what is overdue.',
                   'bi-calendar-week', False, shipment_plan)),
    ('materials', ('Material requirement', 'Every BOM material with its total requirement, supplier and PO status.',
                   'bi-diagram-3', False, material_requirement)),
    ('suppliers', ('Supplier summary', 'POs per supplier: how many, which types, open vs received.',
                   'bi-building', False, supplier_summary)),
    ('po-register', ('Supplier PO register', 'Every supplier PO with PI, LC, ETD, ETA and status.',
                     'bi-truck', False, po_register)),
    ('profitability', ('Costing & margin', 'Cost per piece against FOB and the margin for every style.',
                       'bi-graph-up-arrow', True, profitability)),
    ('material-cost', ('Material cost by type', 'What each order spends on fabric, trims, labels and packing.',
                       'bi-pie-chart', True, material_cost_by_type)),
])


def build(key, filters, admin):
    builder = REPORTS[key][4]
    return builder(filtered_orders(filters)).for_user(admin)


# --------------------------------------------------------------------------
# downloads
# --------------------------------------------------------------------------

def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    return '' if value is None else value


def as_csv(title, table):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([col.label for col in table.columns])
    for row in table.rows:
        writer.writerow([v.isoformat() if isinstance(v, datetime.date) else _plain(v) for v in row])
    response = HttpResponse('﻿' + out.getvalue(), content_type='text/csv; charset=utf-8')  # BOM: Excel reads UTF-8
    response['Content-Disposition'] = f'attachment; filename="{_filename(title)}.csv"'
    return response


XLSX_FORMATS = {'int': '#,##0', 'qty': '#,##0.00', 'money': '#,##0.00', 'rate': '#,##0.0000', 'pct': '0.0"%"',
                'date': 'dd mmm yyyy'}


def as_xlsx(title, table, subtitle=''):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws['A2'].font = Font(color='666666')
    ws.append([])
    ws.append([col.label for col in table.columns])
    header = ws.max_row
    for cell in ws[header]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='10233F')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in table.rows:
        ws.append([_plain(v) for v in row])
    totals = table.totals
    if totals:
        ws.append([_plain(v) for v in totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='F3F6F9')
    for i, col in enumerate(table.columns, start=1):
        letter = get_column_letter(i)
        fmt = XLSX_FORMATS.get(col.kind)
        longest = len(col.label)
        for (cell,) in ws.iter_rows(min_row=header + 1, min_col=i, max_col=i):
            if fmt:
                cell.number_format = fmt
            longest = max(longest, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 50)
    ws.freeze_panes = ws.cell(row=header + 1, column=1)

    out = io.BytesIO()
    wb.save(out)
    response = HttpResponse(
        out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{_filename(title)}.xlsx"'
    return response


def _filename(title):
    slug = ''.join(ch if ch.isalnum() else '-' for ch in title.lower()).strip('-')
    while '--' in slug:
        slug = slug.replace('--', '-')
    return f'orbit-{slug}-{timezone.localdate():%Y%m%d}'
